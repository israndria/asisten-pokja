"""Regression gate resolver snapshot V2 untuk workflow PL."""

from datetime import date, datetime

from docx import Document

import isi_master_data_pl


class _FakeCell:
    def __init__(self):
        self.Value = None
        self.Value2 = None
        self.NumberFormat = None


class _FakeSheet:
    def __init__(self):
        self.cells = {}

    def Range(self, address):
        return self.cells.setdefault(address, _FakeCell())


class _FakeWorkbook:
    def __init__(self):
        self.sheet = _FakeSheet()

    def Worksheets(self, name):
        assert name == "@ Master Data"
        return self.sheet


def test_create_folder_dates_write_native_date_fields():
    wb = _FakeWorkbook()
    logs = []

    result = isi_master_data_pl.set_create_folder_dates(
        wb, date(2026, 9, 5), logs.append
    )

    assert result == date(2026, 9, 5)
    assert wb.sheet.Range("H8").Value == 5
    assert wb.sheet.Range("H9").Value == 9
    assert wb.sheet.Range("H10").Value == 2026
    assert wb.sheet.Range("C21").Value2 == 46270
    assert wb.sheet.Range("C21").NumberFormat == "[$-id-ID]dd mmmm yyyy"
    assert logs and "C21=2026-09-05" in logs[0]


def test_create_folder_date_normalizer_accepts_datetime():
    assert isi_master_data_pl._coerce_create_folder_date(
        datetime(2026, 9, 5, 14, 30)
    ) == date(2026, 9, 5)


def test_master_data_v2_resolver_uses_sibling_procurement_core(monkeypatch, tmp_path):
    code_base = tmp_path / "code"
    asisten_root = code_base / "Asisten_Pokja"
    core_root = code_base / "procurement_core"
    core_root.mkdir(parents=True)
    (core_root / "master_data_v2.py").write_text("# test module\n", encoding="utf-8")

    monkeypatch.setattr(
        isi_master_data_pl,
        "__file__",
        str(asisten_root / "isi_master_data_pl.py"),
    )
    monkeypatch.setenv("POKJA_CODE_ROOT", str(asisten_root))
    monkeypatch.delenv("POKJA_V19_ROOT", raising=False)

    assert isi_master_data_pl._find_master_data_v2_root() == str(core_root.resolve())


def test_master_data_v2_resolver_prefers_configured_v19_root(monkeypatch, tmp_path):
    configured_root = tmp_path / "procurement_core"
    configured_root.mkdir()
    (configured_root / "master_data_v2.py").write_text("# test module\n", encoding="utf-8")

    monkeypatch.setenv("POKJA_CODE_ROOT", str(tmp_path / "wrong"))
    monkeypatch.setenv("POKJA_V19_ROOT", str(configured_root))

    assert isi_master_data_pl._find_master_data_v2_root() == str(configured_root.resolve())


def test_equipment_parser_uses_table_headers_not_filename(tmp_path):
    source = tmp_path / "jln pangkaran.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=6)
    for cell, value in zip(
        table.rows[0].cells,
        ("No", "Nama Peralatan", "Kode Alat", "Kapasitas", "Jumlah (unit)", "Keterangan"),
    ):
        cell.text = value
    for values in (
        ("1.", "Dump Truck", "E08", "3,5 Ton", "1", ""),
        ("2.", "Tandem Roller", "E17", "6-8 T; 74 HP", "1", ""),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    doc.save(source)

    assert isi_master_data_pl._parse_equipment_docx(str(source)) == [
        {"nama": "Dump Truck", "kapasitas": "3,5 Ton", "jumlah": "1 Unit"},
        {"nama": "Tandem Roller", "kapasitas": "6-8 T; 74 HP", "jumlah": "1 Unit"},
    ]


def test_equipment_quantity_normalizes_existing_unit_spelling():
    assert isi_master_data_pl.normalize_equipment_quantity("1") == "1 Unit"
    assert isi_master_data_pl.normalize_equipment_quantity("1 unit") == "1 Unit"
